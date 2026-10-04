from openpyxl import load_workbook


def update_cell(file_path, cell, new_value):
    workbook = load_workbook(file_path)
    sheet = workbook.active

    old_value = sheet[cell].value

    sheet[cell] = new_value

    workbook.save(file_path)
    workbook.close()

    return {
        "cell": cell,
        "old_value": old_value,
        "new_value": new_value
    }


def modify_cell(file_path, cell, command):
    workbook = load_workbook(file_path)
    sheet = workbook.active

    old_value = sheet[cell].value
    operation = command["operation"]

    if operation == "update":

        new_value = command["new_value"]

    elif operation == "increase":

        amount = command["amount"]
        unit = command["unit"]

        if unit == "percent":

            new_value = old_value + (
                old_value * amount / 100
            )

        elif unit == "fixed":

            new_value = old_value + amount

        else:

            raise ValueError(
                f"Unsupported unit: {unit}"
            )

    elif operation == "decrease":

        amount = command["amount"]
        unit = command["unit"]

        if unit == "percent":

            new_value = old_value - (
                old_value * amount / 100
            )

        elif unit == "fixed":

            new_value = old_value - amount

        else:

            raise ValueError(
                f"Unsupported unit: {unit}"
            )

    else:

        raise ValueError(
            f"Operation '{operation}' "
            "is not supported by the cell modifier."
        )

    sheet[cell] = new_value

    workbook.save(file_path)
    workbook.close()

    return {
        "cell": cell,
        "old_value": old_value,
        "new_value": new_value
    }


def add_row(file_path, values):
    workbook = load_workbook(file_path)
    sheet = workbook.active

    headers = {}

    for cell in sheet[1]:

        if cell.value is not None:
            headers[str(cell.value)] = cell.column

    new_row = sheet.max_row + 1

    for column_name, value in values.items():

        if column_name not in headers:

            workbook.close()

            raise ValueError(
                f"Column '{column_name}' "
                "does not exist in Excel."
            )

        column_number = headers[column_name]

        sheet.cell(
            row=new_row,
            column=column_number
        ).value = value

    workbook.save(file_path)
    workbook.close()

    return {
        "row": new_row,
        "values": values
    }


def delete_row(file_path, row):
    """
    Delete exactly one Excel row.

    The row number is resolved by the target-finding
    layer, so this function does not need to understand
    natural language.
    """

    workbook = load_workbook(file_path)
    sheet = workbook.active

    if row < 2 or row > sheet.max_row:

        workbook.close()

        raise ValueError(
            f"Invalid row number: {row}"
        )

    deleted_values = {}

    for cell in sheet[row]:

        header = sheet.cell(
            row=1,
            column=cell.column
        ).value

        if header is not None:
            deleted_values[str(header)] = cell.value

    sheet.delete_rows(row, 1)

    workbook.save(file_path)
    workbook.close()

    return {
        "row": row,
        "deleted_values": deleted_values
    }