from openpyxl import load_workbook


def verify_cell(file_path, cell, expected_value):

    workbook = load_workbook(file_path)
    sheet = workbook.active

    actual_value = sheet[cell].value

    workbook.close()

    verification_passed = (
        actual_value == expected_value
    )

    return {
        "expected_value": expected_value,
        "actual_value": actual_value,
        "passed": verification_passed
    }


def verify_row(file_path, row, expected_values):

    workbook = load_workbook(file_path)
    sheet = workbook.active

    headers = {}

    for cell in sheet[1]:

        if cell.value is not None:
            headers[str(cell.value)] = cell.column

    actual_values = {}

    for column_name, expected_value in expected_values.items():

        if column_name not in headers:

            workbook.close()

            raise ValueError(
                f"Column '{column_name}' "
                "does not exist in Excel."
            )

        column_number = headers[column_name]

        actual_value = sheet.cell(
            row=row,
            column=column_number
        ).value

        actual_values[column_name] = actual_value

    workbook.close()

    verification_passed = True

    for column_name, expected_value in expected_values.items():

        if actual_values[column_name] != expected_value:
            verification_passed = False
            break

    return {
        "row": row,
        "expected_values": expected_values,
        "actual_values": actual_values,
        "passed": verification_passed
    }


def verify_row_deleted(
    file_path,
    target_column,
    target_value
):
    """
    Verify that a row containing the target value
    no longer exists.
    """

    workbook = load_workbook(file_path)
    sheet = workbook.active

    headers = {}

    for cell in sheet[1]:

        if cell.value is not None:
            headers[str(cell.value)] = cell.column

    if target_column not in headers:

        workbook.close()

        raise ValueError(
            f"Column '{target_column}' "
            "does not exist in Excel."
        )

    column_number = headers[target_column]

    remaining_rows = []

    for row in range(2, sheet.max_row + 1):

        value = sheet.cell(
            row=row,
            column=column_number
        ).value

        if str(value).strip().lower() == str(
            target_value
        ).strip().lower():

            remaining_rows.append(row)

    workbook.close()

    return {
        "target_column": target_column,
        "target_value": target_value,
        "remaining_rows": remaining_rows,
        "passed": len(remaining_rows) == 0
    }