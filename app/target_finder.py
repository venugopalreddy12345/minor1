from openpyxl import load_workbook


def find_target(file_path, target_column, target_value, target_field):
    workbook = load_workbook(file_path)
    sheet = workbook.active

    # Find target column
    target_column_number = None

    for cell in sheet[1]:
        if cell.value == target_column:
            target_column_number = cell.column
            break

    if target_column_number is None:
        workbook.close()
        raise ValueError(
            f"Column '{target_column}' not found."
        )

    # Find ALL matching rows
    matching_rows = []

    for row in range(2, sheet.max_row + 1):
        cell_value = sheet.cell(
            row=row,
            column=target_column_number
        ).value

        if cell_value == target_value:
            matching_rows.append(row)

    # No match
    if len(matching_rows) == 0:
        workbook.close()
        raise ValueError(
            f"Value '{target_value}' not found."
        )

    # Multiple matches = ambiguous
    if len(matching_rows) > 1:
        workbook.close()
        raise ValueError(
            f"Ambiguous target: '{target_value}' "
            f"appears in multiple rows: {matching_rows}. "
            f"Please provide more specific information."
        )

    # Exactly one match
    target_row = matching_rows[0]

    # Find field column
    field_column_number = None

    for cell in sheet[1]:
        if cell.value == target_field:
            field_column_number = cell.column
            break

    if field_column_number is None:
        workbook.close()
        raise ValueError(
            f"Field '{target_field}' not found."
        )

    target_cell = sheet.cell(
        row=target_row,
        column=field_column_number
    )

    result = {
        "row": target_row,
        "column": target_cell.column_letter,
        "cell": target_cell.coordinate,
        "current_value": target_cell.value
    }

    workbook.close()

    return result