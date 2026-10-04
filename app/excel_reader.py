from openpyxl import load_workbook


def read_excel_structure(file_path):

    workbook = load_workbook(file_path)

    sheet = workbook.active

    columns = []

    for cell in sheet[1]:
        if cell.value is not None:
            columns.append(str(cell.value))

    structure = {
        "sheet": sheet.title,
        "columns": columns,
        "rows": sheet.max_row,
        "column_count": sheet.max_column
    }

    workbook.close()

    return structure


def find_excel_values(
    file_path,
    target_column=None,
    target_value=None,
    field=None
):

    workbook = load_workbook(file_path)

    sheet = workbook.active

    headers = {}

    for cell in sheet[1]:
        if cell.value is not None:
            headers[str(cell.value)] = cell.column

    # -------------------------------------------------
    # FIND A SPECIFIC PERSON / ROW
    # -------------------------------------------------

    if target_column and target_value:

        if target_column not in headers:
            workbook.close()

            raise ValueError(
                f"Column '{target_column}' not found."
            )

        target_column_number = headers[target_column]

        matching_rows = []

        for row in range(2, sheet.max_row + 1):

            value = sheet.cell(
                row=row,
                column=target_column_number
            ).value

            if str(value).strip().lower() == str(
                target_value
            ).strip().lower():

                matching_rows.append(row)

        if not matching_rows:

            workbook.close()

            raise ValueError(
                f"Value '{target_value}' not found."
            )

        if len(matching_rows) > 1:

            workbook.close()

            raise ValueError(
                f"Ambiguous target: '{target_value}' "
                f"appears in multiple rows: {matching_rows}."
            )

        row_number = matching_rows[0]

        # If no field was supplied, return the entire row.
        if field is None:

            result = {}

            for column_name, column_number in headers.items():

                result[column_name] = sheet.cell(
                    row=row_number,
                    column=column_number
                ).value

            workbook.close()

            return result

        if field not in headers:

            workbook.close()

            raise ValueError(
                f"Field '{field}' not found."
            )

        value = sheet.cell(
            row=row_number,
            column=headers[field]
        ).value

        workbook.close()

        return {
            "field": field,
            "value": value
        }

    # -------------------------------------------------
    # RETURN ENTIRE COLUMN
    # -------------------------------------------------

    if field:

        if field not in headers:

            workbook.close()

            raise ValueError(
                f"Field '{field}' not found."
            )

        values = []

        for row in range(2, sheet.max_row + 1):

            value = sheet.cell(
                row=row,
                column=headers[field]
            ).value

            if value is not None:
                values.append(value)

        workbook.close()

        return {
            "field": field,
            "values": values
        }

    workbook.close()

    raise ValueError(
        "No valid read target was provided."
    )